#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""eBayで売れているキャラの一覧を Excel に出す (2026-09-21).

ルール2 (キャラ軸) の材料。テラピーク台帳 676種類をカタログで引いて、
キャラごとに 売れた枚数 を合算する。仕入上限はキャラ軸では使わない
(会社上限 7万円のみ・値付けは cost-plus) ので、相場は参考値として添えるだけ。

    python iMakHQ/tools/chara_demand_xlsx.py
"""
from __future__ import annotations

import collections
import csv
import os
import re
import sqlite3
import statistics as st
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import market_ledger as M
import catalog_lookup as L

LEDGER = r"C:/dev/iMak_data/hq/market_sold/ledger.csv"
CATALOG = "file:C:/dev/iMak_data/catalog/products.sqlite?mode=ro"
OUT = r"C:/dev/iMak_data/hq/market_sold/chara_demand.xlsx"

# ポケモンの形態 (V / VMAX / ex …) はキャラとして同じ物にまとめる
_SUFFIX = re.compile(r"(VSTAR|VMAX|V-UNION|GX|EX|ex|V)$")


_PROMO_SET = re.compile(r"^(?:[A-Z0-9]+-P|XYP|SMP|BWP|DPP)$", re.I)
_PARA_ID = re.compile(r"(_p\d*|_PARA)(_|$)", re.I)
# ★2026-09-21 ユーザー確定: **ワンピースでは** Alt Art / Parallel / Manga = パラレル。
#   9/20 に markers() から Alt Art を外したのは ドラゴンボール FB05-119 (カード自体が
#   別イラストの SCR で、市場が "Alt Art" と呼んでいた) のため。そちらは変えない。
#   ゲームごとの差は表 (データ) で持つ。if 分岐を増やさない。
_PARA_WORDS = {
    "one_piece_tcg": ("ALT ART", "ALTERNATE", "PARALLEL", "MANGA", "パラレル"),
}
_PARA_WORDS_DEFAULT = ("PARALLEL", "パラレル")


def version_of(product_id, title="", game=""):
    """カードの版 (純関数): 「パラレル」/「プロモ」/「通常」。

    ★product_id の形を先に見る (カタログの値なので確実)。タイトルの語は後ろ盾。
      パラレル: `OP05-119_p1` / `FB08-121_PARA` / ワンピースのタイトルに Alt Art 等
      プロモ  : 弾が `S-P` `M-P` `SV-P` `S8a-P` `XYP` / ワンピースの `P-043`
    """
    pid = (product_id or "").strip()
    t = (title or "").upper()
    words = _PARA_WORDS.get(game, _PARA_WORDS_DEFAULT)
    if _PARA_ID.search(pid) or any(w in t for w in words):
        return "パラレル"
    head = pid.split("_")[0]
    set_part = head.rsplit("-", 1)[0] if "-" in head else ""
    if head.upper().startswith("P-") or _PROMO_SET.match(set_part) or "プロモ" in M.markers(title):
        return "プロモ"
    return "通常"


def chara_of(name):
    """カード名 → キャラ名 (純関数)。「リザードンVSTAR」→「リザードン」。

    ★全角・半角は揃える。カタログに「モンキー・D・ルフィ」と「モンキー・Ｄ・ルフィ」の
      両方があり、揃えないと同じキャラが2つに割れる (2026-09-21)。
    """
    import unicodedata
    n = unicodedata.normalize("NFKC", (name or "").strip())
    return _SUFFIX.sub("", n).strip()


def build(rows, conn):
    """台帳の **出品1行ごと** → キャラごとの集計 (I/O はカタログ引きのみ)。

    ★番号で束ねた需要表 (demand_full.csv) からは作らない。束ねた時点で
      同じ番号の 通常版 と パラレル が1行に混ざり、版が分けられなくなる (2026-09-21)。
    """
    import demand_table_build as D
    agg = collections.defaultdict(lambda: {
        "枚数": 0, "カード": set(), "価格": [], "ゲーム": "", "例": [],
        "版": collections.Counter()})
    cache, miss = {}, 0
    for r in rows:
        title = r.get("タイトル") or ""
        if M.markers(title) & M._OTHER_LANG:     # 英語版などは番号が同じ日本語版と取り違える
            miss += 1
            continue
        nums = D.card_numbers(title)
        if not nums:
            miss += 1
            continue
        num = nums[0]
        try:
            sold = int(float(r.get("売れた数") or 0))
        except ValueError:
            sold = 0
        # ★「引けない」は覚えない。番号だけで覚えると、最初の1件が英語版だった番号は
        #   後の日本語版まで「引けない」になる (世界中のセラーが混ざる台帳で実害)。
        #   引けた結果だけを 番号+弾コード で覚える
        ck = (num, (L.set_code_from_title(title, conn) or "").upper())
        row = cache.get(ck)
        if row is None:
            try:
                row = M.lookup_catalog(M.product_id_candidates(num, title), conn, title)
            except Exception:                                  # noqa: BLE001
                row = None
            if row:
                cache[ck] = row
        if not row or not (row[2] or row[1] or "").strip():
            miss += 1
            continue
        k = chara_of((row[2] or row[1]).strip())
        game = row[3] or ""
        a = agg[k]
        a["枚数"] += sold
        a["カード"].add(row[0])
        a["ゲーム"] = a["ゲーム"] or game
        a["版"][version_of(row[0], title, game)] += sold
        p = D.to_num(r.get("平均落札"))
        if p > 0:
            a["価格"].append(p)
        if len(a["例"]) < 3 and num not in a["例"]:
            a["例"].append(num)
    for a in agg.values():
        a["カード数"] = len(a.pop("カード"))
    return agg, miss


def main():
    rows = list(csv.DictReader(open(LEDGER, encoding="utf-8-sig")))
    conn = sqlite3.connect(CATALOG, uri=True)
    agg, miss = build(rows, conn)

    out = []
    for k, a in agg.items():
        px = sorted(a["価格"])
        out.append({
            "キャラ": k,
            "売れた枚数": a["枚数"],
            "カード種類": a["カード数"],
            "ゲーム": a["ゲーム"],
            "実売中央の中央値(USD)": round(st.median(px), 2) if px else "",
            "実売の最安(USD)": px[0] if px else "",
            "実売の最高(USD)": px[-1] if px else "",
            "うち通常": a["版"].get("通常", 0),
            "うちプロモ": a["版"].get("プロモ", 0),
            "うちパラレル": a["版"].get("パラレル", 0),
            "カード例": " / ".join(a["例"]),
        })
    out.sort(key=lambda r: (-r["売れた枚数"], -r["カード種類"]))

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "キャラ需要"
    cols = list(out[0].keys())
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center")
    for r in out:
        ws.append([r[c] for c in cols])
    widths = [22, 11, 11, 16, 20, 15, 15, 10, 10, 11, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # キャラ×版 (1行 = キャラと版の組)
    ws2 = wb.create_sheet("キャラ×版")
    ws2.append(["キャラ", "版", "売れた枚数", "ゲーム"])
    pairs = []
    for k, a in agg.items():
        for v, n in a["版"].items():
            if n:
                pairs.append((k, v, n, a["ゲーム"]))
    for r in sorted(pairs, key=lambda x: -x[2]):
        ws2.append(list(r))
    for c in range(1, 5):
        cell = ws2.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
    for col, w in zip("ABCD", (22, 10, 11, 16)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    print(f"台帳 {len(rows)}出品 → カタログを引けた {len(rows)-miss} / 引けない {miss}")
    print(f"キャラ {len(out)}種類 / 売れた枚数 合計 {sum(r['売れた枚数'] for r in out)}枚")
    for n in (10, 5, 3, 2):
        print(f"  {n}枚以上: {len([r for r in out if r['売れた枚数'] >= n]):3d}キャラ")
    tot = collections.Counter()
    for a in agg.values():
        tot.update(a["版"])
    print(f"版ごとの売れた枚数: 通常 {tot['通常']} / プロモ {tot['プロモ']} / パラレル {tot['パラレル']}")
    print(f"出力: {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
